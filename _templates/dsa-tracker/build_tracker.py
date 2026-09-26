"""Build the downloadable tracker for the DSA series.

    python _templates/dsa-tracker/build_tracker.py

Reads _data/dsa_problems.yml (the same list the posts render) and writes
assets/downloads/dsa-tracker.xlsx: one row per problem with the columns from the
solving cycle in Part 1, review dates computed from the date solved (1, 3, 7 and
14 days), and a dashboard of cold-solve ratio per topic. Needs openpyxl and PyYAML.
"""
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
TOPICS = yaml.safe_load((ROOT / "_data" / "dsa_problems.yml").read_text(encoding="utf-8"))
OUT = ROOT / "assets" / "downloads" / "dsa-tracker.xlsx"
SITE = "https://rishabh0111.github.io"

FONT = "Arial"
INK = "2E3440"
f_body = Font(name=FONT, size=10, color=INK)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=14, bold=True, color=INK)
f_note = Font(name=FONT, size=9, italic=True, color="5E6A83")
f_link = Font(name=FONT, size=10, color="1D6F87", underline="single")
fill_head = PatternFill("solid", fgColor="3B4252")
fill_input_head = PatternFill("solid", fgColor="1E768F")
fill_input = PatternFill("solid", fgColor="FFF9DB")  # cells you fill in
thin = Side(style="thin", color="D8DEE9")
border = Border(bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
top = Alignment(vertical="top")
center = Alignment(horizontal="center", vertical="top")

DIFF_FILL = {"Easy": ("D1FAE5", "065F46"), "Medium": ("FEF3C7", "92400E"), "Hard": ("FEE2E2", "991B1B")}

wb = Workbook()

# ── Tracker ────────────────────────────────────────────────────────────────
ws = wb.active
ws.title = "Tracker"
cols = [
    # header, width, is_input
    ("#", 5, False),
    ("Part", 6, False),
    ("Topic", 22, False),
    ("Problem", 40, False),
    ("Difficulty", 11, False),
    ("Solved", 14, True),
    ("Date solved", 13, True),
    ("Key sentence", 50, True),
    ("Pattern", 22, True),
    ("Time", 12, True),
    ("Space", 10, True),
    ("Edge case missed", 28, True),
    ("Reviews passed", 10, True),
    ("Next review", 13, False),
    ("Status", 12, False),
]
for i, (name, width, is_input) in enumerate(cols, 1):
    c = ws.cell(row=1, column=i, value=name)
    c.font = f_head
    c.fill = fill_input_head if is_input else fill_head
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[1].height = 30

r = 2
for t in TOPICS:
    for p in t["problems"]:
        ws.cell(row=r, column=1, value=r - 1).alignment = center
        ws.cell(row=r, column=2, value=t["part"]).alignment = center
        ws.cell(row=r, column=3, value=t["topic"])
        link = ws.cell(row=r, column=4, value=p["name"])
        link.hyperlink = p["url"]
        link.font = f_link
        d = ws.cell(row=r, column=5, value=p["difficulty"])
        bg, fg = DIFF_FILL[p["difficulty"]]
        d.fill = PatternFill("solid", fgColor=bg)
        d.font = Font(name=FONT, size=10, color=fg, bold=True)
        d.alignment = center
        # Next review: 1, 3, 7 and 14 days after the date solved, by reviews passed.
        ws.cell(row=r, column=14, value=f'=IF(G{r}="","",IF(N(M{r})>=4,"done",G{r}+CHOOSE(N(M{r})+1,1,3,7,14)))')
        ws.cell(row=r, column=15, value=f'=IF(G{r}="","not started",IF(N{r}="done","done",IF(N{r}<=TODAY(),"due","scheduled")))')
        for col in range(1, 16):
            c = ws.cell(row=r, column=col)
            if c.font != f_link and col != 5:
                c.font = f_body
            c.border = border
            if col not in (1, 2, 5):
                c.alignment = wrap if col in (8, 12) else top
            if cols[col - 1][2]:
                c.fill = fill_input
        ws.cell(row=r, column=7).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=13).alignment = center
        ws.cell(row=r, column=14).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=14).alignment = center
        ws.cell(row=r, column=15).alignment = center
        r += 1
last = r - 1

dv_solved = DataValidation(type="list", formula1='"Cold,Needed video"', allow_blank=True,
                           promptTitle="Solved", prompt="Cold: solved with no help. Needed video: needed a walkthrough or the chat's approach.")
dv_reviews = DataValidation(type="whole", operator="between", formula1="0", formula2="4", allow_blank=True,
                            promptTitle="Reviews passed", prompt="0 to 4. A failed review: set this back to 0 and Date solved to today.")
dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                         promptTitle="Date solved", prompt="The day you first solved it (yyyy-mm-dd).")
for dv, col in ((dv_solved, "F"), (dv_reviews, "M"), (dv_date, "G")):
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{last}")

rng = f"A2:O{last}"
ws.conditional_formatting.add(f"O2:O{last}", FormulaRule(formula=['$O2="due"'], fill=PatternFill("solid", fgColor="FEE2E2"), font=Font(name=FONT, bold=True, color="991B1B")))
ws.conditional_formatting.add(f"O2:O{last}", FormulaRule(formula=['$O2="done"'], fill=PatternFill("solid", fgColor="D1FAE5"), font=Font(name=FONT, color="065F46")))
ws.conditional_formatting.add(f"F2:F{last}", FormulaRule(formula=['$F2="Cold"'], fill=PatternFill("solid", fgColor="D1FAE5"), font=Font(name=FONT, color="065F46")))
ws.conditional_formatting.add(f"F2:F{last}", FormulaRule(formula=['$F2="Needed video"'], fill=PatternFill("solid", fgColor="FEF3C7"), font=Font(name=FONT, color="92400E")))
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:O{last}"

# ── Dashboard ──────────────────────────────────────────────────────────────
db = wb.create_sheet("Dashboard")
db["A1"] = "Progress"
db["A1"].font = f_title
db["A2"] = "Everything here is computed from the Tracker sheet. Cold ratio = cold solves / solved; rising means the method is working, flat on one topic means revisit its lesson."
db["A2"].font = f_note
db.merge_cells("A2:J2")
db["A2"].alignment = Alignment(wrap_text=True, vertical="top")
db.row_dimensions[2].height = 28

heads = ["Part", "Topic", "Problems", "Easy", "Medium", "Hard", "Solved", "Cold", "Needed video", "Cold ratio", "Due now"]
for i, h in enumerate(heads, 1):
    c = db.cell(row=4, column=i, value=h)
    c.font = f_head
    c.fill = fill_head
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
T = f"Tracker!$C$2:$C${last}"
row = 5
for t in TOPICS:
    db.cell(row=row, column=1, value=t["part"]).alignment = center
    db.cell(row=row, column=2, value=t["topic"])
    b = f"$B{row}"
    db.cell(row=row, column=3, value=f"=COUNTIF({T},{b})")
    for j, lvl in enumerate(("Easy", "Medium", "Hard")):
        db.cell(row=row, column=4 + j, value=f'=COUNTIFS({T},{b},Tracker!$E$2:$E${last},"{lvl}")')
    db.cell(row=row, column=7, value=f'=COUNTIFS({T},{b},Tracker!$F$2:$F${last},"<>")')
    db.cell(row=row, column=8, value=f'=COUNTIFS({T},{b},Tracker!$F$2:$F${last},"Cold")')
    db.cell(row=row, column=9, value=f'=COUNTIFS({T},{b},Tracker!$F$2:$F${last},"Needed video")')
    db.cell(row=row, column=10, value=f'=IF(G{row}=0,"",H{row}/G{row})')
    db.cell(row=row, column=11, value=f'=COUNTIFS({T},{b},Tracker!$O$2:$O${last},"due")')
    for col in range(1, 12):
        c = db.cell(row=row, column=col)
        c.font = f_body
        c.border = border
        if col >= 3:
            c.alignment = center
    db.cell(row=row, column=10).number_format = "0%"
    row += 1
tot = row
db.cell(row=tot, column=2, value="All topics").font = Font(name=FONT, size=10, bold=True, color=INK)
for col in (3, 4, 5, 6, 7, 8, 9, 11):
    L = get_column_letter(col)
    c = db.cell(row=tot, column=col, value=f"=SUM({L}5:{L}{tot - 1})")
    c.font = Font(name=FONT, size=10, bold=True, color=INK)
    c.alignment = center
c = db.cell(row=tot, column=10, value=f'=IF(G{tot}=0,"",H{tot}/G{tot})')
c.number_format = "0%"
c.font = Font(name=FONT, size=10, bold=True, color=INK)
c.alignment = center
for col in range(1, 12):
    db.cell(row=tot, column=col).border = Border(top=Side(style="medium", color=INK))
db.conditional_formatting.add(f"K5:K{tot}", FormulaRule(formula=["K5>0"], fill=PatternFill("solid", fgColor="FEE2E2"), font=Font(name=FONT, bold=True, color="991B1B")))
for col, w in zip("ABCDEFGHIJK", (6, 24, 10, 8, 9, 8, 9, 8, 12, 11, 10)):
    db.column_dimensions[col].width = w
db.freeze_panes = "A5"

# ── How to use ─────────────────────────────────────────────────────────────
hw = wb.create_sheet("How to use")
hw.column_dimensions["A"].width = 22
hw.column_dimensions["B"].width = 95
hw["A1"] = "How to use this tracker"
hw["A1"].font = f_title
lines = [
    ("Where it's from", f"The solving cycle in Part 1 of Data Structures & Algorithms, Pattern by Pattern: {SITE}/blogs/dsa-the-method/"),
    ("You fill in", "Only the cells shaded pale yellow on the Tracker sheet (teal headers): Solved, Date solved, Key sentence, Pattern, Time, Space, Edge case missed, Reviews passed. Everything else is filled or computed."),
    ("Solved", "Cold if you solved it after the 15-minute struggle with no help. Needed video if you needed a walkthrough or the chat's approach. Cold solves are the metric."),
    ("Key sentence", "One line you could rebuild the whole solution from. If you can't rebuild it from the line, the line was too vague: rewrite it."),
    ("Reviews", "Re-solve 1, 3, 7 and 14 days after the date solved. Next review and Status work it out. A review passes if you can state the pattern, the key idea and the complexity fluently in about two minutes; add 1 to Reviews passed."),
    ("A failed review", "Struggle 5 minutes, read the key sentence, re-watch only as a last resort. Then set Date solved to today and Reviews passed to 0: the problem starts again from day 1."),
    ("Status", "not started · scheduled · due (review today or overdue, shown red) · done (all four reviews passed)."),
    ("Dashboard", "Per-topic progress and cold ratio. Filter the Tracker's Status column to 'due' for today's reviews."),
    ("Google Sheets", "File > Import > Upload this file. The formulas and dropdowns carry over."),
]
r = 3
for k, v in lines:
    hw.cell(row=r, column=1, value=k).font = Font(name=FONT, size=10, bold=True, color=INK)
    c = hw.cell(row=r, column=2, value=v)
    c.font = f_body
    c.alignment = Alignment(wrap_text=True, vertical="top")
    hw.cell(row=r, column=1).alignment = top
    r += 1
r += 1
hw.cell(row=r, column=1, value="Example row").font = Font(name=FONT, size=10, bold=True, color=INK)
hw.cell(row=r, column=2, value="How a filled-in Tracker row looks (a made-up problem, not one from the list):").font = f_note
r += 1
example = [
    ("Problem", "Longest run of equal neighbours"), ("Solved", "Cold"), ("Date solved", "2024-02-10"),
    ("Key sentence", "One pass, keep the current run length; reset to 1 when a[i] != a[i-1]; answer is the max run seen."),
    ("Pattern", "Linear scan with a running state"), ("Time", "O(n)"), ("Space", "O(1)"),
    ("Edge case missed", "Empty input should return 0, not 1"), ("Reviews passed", "2  → next review is 7 days after 2024-02-10"),
]
for k, v in example:
    hw.cell(row=r, column=1, value=k).font = f_body
    c = hw.cell(row=r, column=2, value=v)
    c.font = f_body
    c.fill = fill_input
    r += 1

for s in (ws, db, hw):
    s.sheet_view.showGridLines = s is ws

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT, last - 1, "problems")
